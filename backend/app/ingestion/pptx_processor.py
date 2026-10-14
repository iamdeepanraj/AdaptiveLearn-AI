from pptx import Presentation
from typing import List, Dict, Any
import logging
import os

logger = logging.getLogger(__name__)

class PPTXProcessor:
    @staticmethod
    def process_pptx(file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts text, slide numbers, notes, and shape descriptions from PowerPoint files.
        Returns a list of slide dicts:
        [{
            'slide_number': int (1-indexed),
            'title': str,
            'text': str,
            'notes': str,
            'figures': List[Dict[str, Any]],
            'word_count': int
        }]
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"PPTX file not found at: {file_path}")

        extracted_slides = []
        prs = Presentation(file_path)

        for slide_index, slide in enumerate(prs.slides):
            slide_number = slide_index + 1
            slide_title = ""
            text_runs = []
            figures = []

            for shape_idx, shape in enumerate(slide.shapes):
                # Title shape
                if shape == slide.shapes.title and shape.has_text_frame:
                    slide_title = shape.text.strip()
                    text_runs.append(f"Title: {slide_title}")
                elif shape.has_text_frame:
                    shape_text = shape.text.strip()
                    if shape_text:
                        text_runs.append(shape_text)
                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if row_vals:
                            table_rows.append(" | ".join(row_vals))
                    if table_rows:
                        text_runs.append("Table Data:\n" + "\n".join(table_rows))
                elif shape.shape_type == 13: # Picture
                    figures.append({
                        "figure_id": f"slide{slide_number}_img{shape_idx + 1}",
                        "description": f"Slide {slide_number} embedded illustration / chart: {shape.name}"
                    })

            # Check for speaker notes
            notes_text = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes_text = slide.notes_slide.notes_text_frame.text.strip()

            full_text = "\n\n".join(text_runs)
            if notes_text:
                full_text += f"\n\nSpeaker Notes: {notes_text}"

            extracted_slides.append({
                "slide_number": slide_number,
                "title": slide_title,
                "text": full_text,
                "notes": notes_text,
                "figures": figures,
                "word_count": len(full_text.split())
            })

        logger.info(f"Successfully processed PPTX: {file_path} with {len(extracted_slides)} slides.")
        return extracted_slides
