import os
import fitz
from pptx import Presentation
from pptx.util import Inches, Pt
from backend.app.ingestion.pdf_processor import PDFProcessor
from backend.app.ingestion.pptx_processor import PPTXProcessor
from backend.app.ingestion.chunker import SemanticChunker
from backend.app.models.document import SourceType

def test_pdf_extraction_preserves_page_numbers(tmp_path):
    # Generate a real 2-page PDF using PyMuPDF
    pdf_path = str(tmp_path / "test_sample.pdf")
    doc = fitz.open()
    
    # Page 1
    p1 = doc.new_page()
    p1.insert_text((50, 72), "Chapter 1: Neural Networks and Loss Gradients. Supervised models learn parameters through empirical error reduction.")
    
    # Page 2
    p2 = doc.new_page()
    p2.insert_text((50, 72), "Chapter 2: Backpropagation algorithm uses recursive chain rule calculations to compute derivatives.")
    doc.save(pdf_path)
    doc.close()

    # Process with PDFProcessor
    results = PDFProcessor.process_pdf(pdf_path)
    assert len(results) == 2
    assert results[0]["page_number"] == 1
    assert "Supervised models" in results[0]["text"]
    assert results[1]["page_number"] == 2
    assert "Backpropagation" in results[1]["text"]

def test_pptx_extraction_preserves_slide_numbers(tmp_path):
    # Generate a real 2-slide PPTX using python-pptx
    pptx_path = str(tmp_path / "test_sample.pptx")
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    
    # Slide 1
    s1 = prs.slides.add_slide(blank_layout)
    txBox = s1.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tf = txBox.text_frame
    tf.text = "Slide 1: Scaled Dot-Product Attention Architecture"

    # Slide 2
    s2 = prs.slides.add_slide(blank_layout)
    txBox2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    tf2 = txBox2.text_frame
    tf2.text = "Slide 2: Multi-Head Projection and Subspace Attention"
    prs.save(pptx_path)

    # Process with PPTXProcessor
    slides = PPTXProcessor.process_pptx(pptx_path)
    assert len(slides) == 2
    assert slides[0]["slide_number"] == 1
    assert "Attention Architecture" in slides[0]["text"]
    assert slides[1]["slide_number"] == 2
    assert "Multi-Head" in slides[1]["text"]

def test_semantic_chunker_preserves_locations():
    items = [
        {"page_number": 3, "text": "Gradient Descent optimization minimizes loss. " * 30, "figures": []},
        {"page_number": 4, "text": "Stochastic updates reduce computational burden per step. " * 30, "figures": []}
    ]
    chunks = SemanticChunker.chunk_document_content(items, source_type=SourceType.PDF, max_words_per_chunk=50)
    assert len(chunks) >= 2
    for c in chunks:
        assert c["page_number"] in [3, 4]
        assert c["source_type"] == SourceType.PDF
        assert len(c["content"]) > 10
